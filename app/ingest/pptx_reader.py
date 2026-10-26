r"""PowerPoint (.pptx) -> `SourceDocument`.

Shapes are visited per slide in reading order -- sorted by (top, left), with
groups expanded in place -- because a deck's XML order is z-order, which has
nothing to do with how a human reads the slide.

Text frames give paragraphs; python-pptx renders an `<a:br>` as "\v", so the
shared visual-line split turns a multi-line option block inside one paragraph
into one `SourceLine` per option, exactly as for Word soft breaks.

`SourceLine.page` is the 1-based slide number and `block_id` counts shapes, so
all lines of one text frame stay grouped.
"""

from __future__ import annotations

import logging
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..models import DocKind, LineKind, SourceDocument
from . import _DocBuilder, _looks_like_question_header

log = logging.getLogger(__name__)

_GROUP_TYPES = frozenset({MSO_SHAPE_TYPE.GROUP})
_PICTURE_TYPES = frozenset({MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE})
#: Guard against a pathologically nested group tree.
_MAX_GROUP_DEPTH = 8


def read_pptx(path: str | Path) -> SourceDocument:
    """Flatten a deck into ordered lines plus its pictures and tables."""
    p = Path(path)
    prs = Presentation(str(p))
    builder = _DocBuilder(str(p), DocKind.PPTX)
    block = -1
    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape_no, shape in enumerate(_ordered_shapes(slide.shapes, 0)):
            block += 1
            _read_shape(shape, slide_no, shape_no, block, builder)
    return builder.finish()


# --------------------------------------------------------------------------- #
# Shape ordering
# --------------------------------------------------------------------------- #


def _position(shape: object) -> tuple[int, int]:
    """(top, left) in EMU; a placeholder inheriting its position sorts first."""
    top = getattr(shape, "top", None)
    left = getattr(shape, "left", None)
    return (int(top) if top is not None else 0, int(left) if left is not None else 0)


def _ordered_shapes(shapes: object, depth: int) -> list[object]:
    """Reading order, with group members spliced in at the group's position."""
    try:
        members = list(shapes)  # type: ignore[call-overload]
    except (AttributeError, TypeError):
        return []
    ordered: list[object] = []
    for shape in sorted(members, key=_position):
        if shape.shape_type in _GROUP_TYPES and depth < _MAX_GROUP_DEPTH:
            ordered.extend(_ordered_shapes(getattr(shape, "shapes", ()), depth + 1))
        else:
            ordered.append(shape)
    return ordered


# --------------------------------------------------------------------------- #
# Per-shape extraction
# --------------------------------------------------------------------------- #


def _read_shape(
    shape: object,
    slide_no: int,
    shape_no: int,
    block: int,
    builder: _DocBuilder,
) -> None:
    """Emit whatever content one shape carries: text, a table or a picture."""
    if getattr(shape, "has_table", False):
        _read_table(shape, slide_no, shape_no, block, builder)
        return
    if shape.shape_type in _PICTURE_TYPES or hasattr(shape, "image"):
        _read_picture(shape, slide_no, builder)
        # A picture placeholder can still hold a caption text frame, so fall
        # through to the text handling below instead of returning.
    if getattr(shape, "has_text_frame", False):
        _read_text_frame(shape, slide_no, shape_no, block, builder)


def _read_text_frame(
    shape: object,
    slide_no: int,
    shape_no: int,
    block: int,
    builder: _DocBuilder,
) -> None:
    for para_no, para in enumerate(shape.text_frame.paragraphs):  # type: ignore
        try:
            text = para.text
        except (AttributeError, TypeError, ValueError) as exc:
            builder.warn(
                f"slide {slide_no}: paragraph text unreadable ({exc})"
            )
            text = ""
        builder.add_text_lines(
            text,
            kind=LineKind.PARAGRAPH,
            block_id=block,
            ref=f"sld:{slide_no}/sh:{shape_no}/p:{para_no}",
            page=slide_no,
        )


def _read_table(
    shape: object,
    slide_no: int,
    shape_no: int,
    block: int,
    builder: _DocBuilder,
) -> None:
    """A data grid goes to `tables`; a single-column question box is flattened."""
    table = shape.table  # type: ignore[attr-defined]
    grid = [[cell.text for cell in row.cells] for row in table.rows]
    if not grid:
        return
    n_cols = max((len(row) for row in grid), default=0)
    is_box = (
        n_cols == 1
        and len(grid) >= 3
        and _looks_like_question_header(grid[0][0])
    )
    if not is_box:
        rows = [tuple(r) for r in grid if any(c.strip() for c in r)]
        builder.add_table(rows, page=slide_no)
        return
    # Same reasoning as in docx_reader: a one-column box is a text container,
    # so the parsers must see its lines rather than an opaque grid.
    for r, row in enumerate(grid):
        for c, cell_text in enumerate(row):
            builder.add_text_lines(
                cell_text,
                kind=LineKind.CELL,
                block_id=block,
                ref=f"sld:{slide_no}/sh:{shape_no}/tbl/r:{r}/c:{c}",
                page=slide_no,
            )


def _read_picture(shape: object, slide_no: int, builder: _DocBuilder) -> None:
    """Record a picture's original bytes, anchored to the text above it."""
    try:
        image = shape.image  # type: ignore[attr-defined]
        blob = image.blob
        width, height = image.size
        ext = image.ext
    except (AttributeError, ValueError, KeyError, OSError) as exc:
        # A linked (not embedded) picture has no blob; nothing to recover.
        builder.warn(f"slide {slide_no}: picture unreadable ({exc}); skipped")
        return
    builder.add_image(blob, int(width), int(height), str(ext), page=slide_no)
